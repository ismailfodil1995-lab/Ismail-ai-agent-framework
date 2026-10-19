import os
from docx import Document
from tools import register_tool


@register_tool(
    name="read_docx_file",
    description="استخراج وقراءة المحتوى النصي من ملفات Word (.docx)."
)
def read_docx_file(file_path: str = "", **kwargs) -> str:
    """
    قراءة المحتوى النصي من ملف Word (.docx) مع معالجة الأخطاء والتوافق التام.
    """
    for key in ["file_path", "filename", "file", "path", "doc_path"]:
        if key in kwargs and kwargs[key]:
            file_path = kwargs[key]
            break

    if not file_path:
        return "Error: No DOCX file path provided."

    if not os.path.exists(file_path):
        return f"Error: File '{file_path}' not found."

    try:
        doc = Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # قراءة الجداول إن وجدت
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    tables_text.append(" | ".join(row_cells))

        full_text = "\n".join(paragraphs)
        if tables_text:
            full_text += "\n\n--- Tables ---\n" + "\n".join(tables_text)

        if not full_text.strip():
            return f"Warning: DOCX file '{file_path}' is empty."

        # تقييد الحجم لمنع استنزاف السياق
        preview = full_text[:4000]
        return f"SUCCESS: Read DOCX '{file_path}':\n{preview}"

    except Exception as e:
        return f"Error reading DOCX file: {str(e)}"
