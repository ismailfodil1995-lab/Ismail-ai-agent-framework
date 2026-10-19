"""
Test Suite: اختبار شامل للمهارات الخمس الجديدة وتأكيد استقرار الـ 23 أداة
"""
import sys
import os
import docx

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')

from brain import AVAILABLE_TOOLS
from services.tool_registry_bridge import tool_bridge
from services.guardrails import guardrails
from executor import executor


def test_tool_count_and_schemas():
    print("\n--- [1/6] التحقق من اكتمال الـ 23 أداة ومخططاتها ---")
    assert len(AVAILABLE_TOOLS) >= 23, f"Expected at least 23 tools, got {len(AVAILABLE_TOOLS)}"
    all_meta = tool_bridge.get_all_tools_metadata()
    assert len(all_meta) >= 23, f"Expected at least 23 tools in bridge, got {len(all_meta)}"
    print(f"✅ تم تأكيد تسجيل {len(AVAILABLE_TOOLS)} أداة في AVAILABLE_TOOLS و ToolRegistryBridge بنجاح.")


def test_docx_tool():
    print("\n--- [2/6] اختبار أداة قراءة ملفات Word (read_docx_file) ---")
    # إنشاء ملف DOCX تجريبي
    doc_path = "sample_test_doc.docx"
    doc = docx.Document()
    doc.add_heading("AI Automation Brief", level=1)
    doc.add_paragraph("This is a verified test document created for the agent framework.")
    doc.save(doc_path)

    cmd = {"tool": "read_docx_file", "args": {"file_path": doc_path}}
    res = executor.execute([cmd])
    print(f"نتيجة قراءة DOCX: {res[:120]}...")
    assert "SUCCESS" in res
    assert "AI Automation Brief" in res

    if os.path.exists(doc_path):
        os.remove(doc_path)
    print("✅ أداة read_docx_file تعمل بنجاح تام.")


def test_email_sender_and_guardrail():
    print("\n--- [3/6] اختبار أداة البريد الإلكتروني (send_email_report) وحماية Guardrails ---")
    assert guardrails.is_critical("send_email_report") is True
    print("✅ تم التحقق من أن send_email_report محمية كأداة حرجة (CRITICAL_TOOL).")

    cmd = {
        "tool": "send_email_report",
        "args": {
            "recipient_email": "client@example.com",
            "subject": "Q3 Enterprise Report",
            "body": "Your requested automation report is ready."
        }
    }
    res = executor.execute([cmd])
    print(f"نتيجة إرسال البريد: {res}")
    assert "SUCCESS" in res
    assert "client@example.com" in res
    print("✅ أداة send_email_report أرسلت وسجلت الإرسال بنجاح.")


def test_json_extractor():
    print("\n--- [4/6] اختبار أداة استخراج الـ JSON (extract_structured_json) ---")
    unstructured_text = "Client Name: ACME Corp, Estimated Budget: $12000, Start Date: Next Monday"
    cmd = {
        "tool": "extract_structured_json",
        "args": {
            "text_data": unstructured_text,
            "schema_keys": ["Client Name", "Estimated Budget", "Start Date"]
        }
    }
    res = executor.execute([cmd])
    print(f"نتيجة استخراج JSON:\n{res}")
    assert "SUCCESS" in res
    assert "ACME Corp" in res
    print("✅ أداة extract_structured_json قامت بهيكلة البيانات بدقة.")


def test_db_tool():
    print("\n--- [5/6] اختبار أداة قاعدة البيانات SQLite (query_local_db) ---")
    db_test = "test_agent.db"

    # 1. إدخال سجل عميل
    insert_cmd = {
        "tool": "query_local_db",
        "args": {
            "query": "INSERT INTO leads (client_name, score, category, notes) VALUES (?, ?, ?, ?)",
            "params": ["Enterprise Client X", 95, "High Quality", "Urgent onboarding requested"],
            "db_name": db_test
        }
    }
    res_insert = executor.execute([insert_cmd])
    assert "SUCCESS" in res_insert
    print(f"نتيجة الإدخال: {res_insert}")

    # 2. استعلام السجلات
    select_cmd = {
        "tool": "query_local_db",
        "args": {
            "query": "SELECT client_name, score, category FROM leads WHERE score > 80",
            "db_name": db_test
        }
    }
    res_select = executor.execute([select_cmd])
    print(f"نتيجة الاستعلام:\n{res_select}")
    assert "SUCCESS" in res_select
    assert "Enterprise Client X" in res_select

    if os.path.exists(db_test):
        os.remove(db_test)
    print("✅ أداة query_local_db نفذت عمليات الكتابة والقراءة الجدولية بنجاح.")


def test_dag_compiler_tool():
    print("\n--- [6/6] اختبار محول ومجمع الـ DAG الديناميكي (compile_dag_workflow) ---")
    cmd = {
        "tool": "compile_dag_workflow",
        "args": {
            "goal": "Client interested in AI Agent, high budget, start now"
        }
    }
    res = executor.execute([cmd])
    print(f"معاينة تنفيذ مسار الـ DAG:\n{res[:200]}...")
    assert "SUCCESS" in res or "COMPLETED" in res
    print("✅ أداة compile_dag_workflow أنشأت ونفذت مسار الـ DAG ديناميكياً.")


if __name__ == "__main__":
    print("==================================================")
    print("🚀 بدء اختبارات المهارات الـ 23 في الـ Framework...")
    print("==================================================")
    test_tool_count_and_schemas()
    test_docx_tool()
    test_email_sender_and_guardrail()
    test_json_extractor()
    test_db_tool()
    test_dag_compiler_tool()
    print("\n" + "=" * 50)
    print("🎉 تم اجتياز جميع اختبارات المهارات الجديدة بنجاح تام (6/6)!")
    print("==================================================")
